#!/usr/bin/env python3
"""Build giant-axonal-neuropathy.pptx — an editable PowerPoint version of the HTML deck.

Usage:  python3 build_pptx.py
Requires: pip install python-pptx pillow

Fonts are limited to Georgia (headings) and Calibri (body), which ship with
Windows/PowerPoint, so the deck renders identically on any PC.
"""
from __future__ import annotations

import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
ASSETS = HERE / "assets"
OUT = HERE / "giant-axonal-neuropathy.pptx"

# ---------------------------------------------------------------- theme
PAPER = RGBColor(0xF6, 0xF3, 0xEC)
INK = RGBColor(0x14, 0x21, 0x3D)
INK_SOFT = RGBColor(0x4A, 0x55, 0x70)
NAVY = RGBColor(0x0F, 0x1B, 0x33)
NAVY_2 = RGBColor(0x18, 0x27, 0x47)
NAVY_GLOW = RGBColor(0x12, 0x30, 0x45)
TEAL = RGBColor(0x1A, 0x9C, 0x8E)
TEAL_SOFT = RGBColor(0xDD, 0xEF, 0xEC)
TEAL_LIGHT = RGBColor(0x7F, 0xE0, 0xD4)
CORAL = RGBColor(0xE4, 0x57, 0x2E)
GOLD = RGBColor(0xE9, 0xB4, 0x4C)
GOLD_DARK = RGBColor(0xB7, 0x80, 0x1E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xDC, 0xD8, 0xCE)
LINE_DARK = RGBColor(0x2A, 0x3A, 0x5C)
PAPER_DIM = RGBColor(0xB8, 0xBC, 0xC8)

HEAD_FONT = "Georgia"
BODY_FONT = "Calibri"

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)

TOKEN = re.compile(r"(\*\*.+?\*\*|__.+?__)")


# ---------------------------------------------------------------- helpers
def set_autofit(text_frame):
    """Ask PowerPoint to shrink text on overflow as a safety net."""
    body = text_frame._txBody.find(qn("a:bodyPr"))
    for child in list(body):
        if child.tag in (qn("a:normAutofit"), qn("a:spAutoFit"), qn("a:noAutofit")):
            body.remove(child)
    body.append(body.makeelement(qn("a:normAutofit"), {}))


def add_runs(paragraph, text, size, color, font=BODY_FONT, bold=False, italic=False):
    """Append runs to a paragraph, honouring **bold** and __italic__ markup."""
    for part in TOKEN.split(text):
        if not part:
            continue
        run = paragraph.add_run()
        r_bold, r_italic, r_font = bold, italic, font
        if part.startswith("**"):
            part, r_bold = part[2:-2], True
        elif part.startswith("__"):
            part, r_italic, r_font = part[2:-2], True, HEAD_FONT
        run.text = part
        f = run.font
        f.name = r_font
        f.size = Pt(size)
        f.bold = r_bold
        f.italic = r_italic
        f.color.rgb = color


def text_box(slide, x, y, w, h, paragraphs, *, size=13, color=INK, font=BODY_FONT,
             bold=False, italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             line_spacing=1.1, space_after=4, autofit=True, margin=0.0):
    """paragraphs: list of str or dict(text=..., size=, color=, bold=, font=, space_after=)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    if autofit:
        set_autofit(tf)
    for i, item in enumerate(paragraphs):
        spec = item if isinstance(item, dict) else {"text": item}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = spec.get("align", align)
        p.line_spacing = spec.get("line_spacing", line_spacing)
        p.space_after = Pt(spec.get("space_after", space_after))
        add_runs(p, spec["text"], spec.get("size", size), spec.get("color", color),
                 spec.get("font", font), spec.get("bold", bold), spec.get("italic", italic))
    return box


def bullets(slide, x, y, w, h, items, *, size=13, color=INK_SOFT, dot=TEAL, gap=7, line_spacing=1.12):
    """Bullet list: coloured dot glyph + rich text. Each item may be str or (str, dot_color)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0))
    set_autofit(tf)
    for i, item in enumerate(items):
        text, dcol = (item if isinstance(item, tuple) else (item, dot))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(gap)
        # hanging indent so wrapped lines align with the text, not the dot
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.28)))
        pPr.set("indent", str(-Inches(0.28)))
        r = p.add_run()
        r.text = "●\u2002"
        r.font.name = BODY_FONT
        r.font.size = Pt(size * 0.7)
        r.font.color.rgb = dcol
        add_runs(p, text, size, color)
    return box


def rect(slide, x, y, w, h, fill, *, line=None, rounded=True, radius=0.06, shadow=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    if not shadow:
        s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def oval(slide, x, y, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def hline(slide, x, y, w, color, weight=1.0):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    return ln


def picture_fit(slide, path, x, y, w, h, *, align="center"):
    """Place an image scaled to fit inside the box, preserving aspect ratio."""
    with Image.open(path) as im:
        iw, ih = im.size
    box_ratio, img_ratio = w / h, iw / ih
    if img_ratio > box_ratio:
        pw, ph = w, w / img_ratio
    else:
        ph, pw = h, h * img_ratio
    px = x + (w - pw) / 2 if align == "center" else x
    py = y + (h - ph) / 2
    pic = slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    return pic, (px, py, pw, ph)


def caption(slide, x, y, w, label, text, *, dark=False):
    col_label = PAPER if dark else INK
    col_text = PAPER_DIM if dark else INK_SOFT
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.35))
    tf = box.text_frame
    tf.word_wrap = True
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0))
    p = tf.paragraphs[0]
    add_runs(p, label.upper() + "  ", 8.5, col_label, bold=True)
    add_runs(p, text, 9, col_text)
    return box


def kicker(slide, x, y, w, text, color=TEAL):
    hline(slide, x, y + 0.11, 0.35, color, 1.5)
    return text_box(slide, x + 0.45, y, w, 0.3, [text.upper()], size=9, color=color, bold=True, autofit=False)


def heading(slide, x, y, w, text, *, dark=False, size=30):
    return text_box(slide, x, y, w, 0.75, [text], size=size, color=PAPER if dark else INK,
                    font=HEAD_FONT, autofit=False, line_spacing=1.0)


def chrome(slide, index, total, left_label, *, dark=False, right_label="Giant axonal neuropathy"):
    """Top bar labels and bottom-right slide counter, echoing the HTML deck."""
    col = PAPER_DIM if dark else INK_SOFT
    text_box(slide, 0.6, 0.28, 6, 0.25, [left_label.upper()], size=8, color=col, bold=True, autofit=False)
    text_box(slide, 6.7, 0.28, 6.0, 0.25, [right_label.upper()], size=8, color=col, bold=True,
             align=PP_ALIGN.RIGHT, autofit=False)
    text_box(slide, 11.2, 7.0, 1.5, 0.25, [f"{index:02d} / {total:02d}"], size=8, color=col, bold=True,
             align=PP_ALIGN.RIGHT, autofit=False)
    # progress bar
    rect(slide, 0, 0, 13.333 * index / total, 0.05, TEAL, rounded=False)


def stat_card(slide, x, y, w, h, value, label, *, value_color=INK, dark=False, value_size=26):
    fill = NAVY_2 if dark else WHITE
    rect(slide, x, y, w, h, fill, line=None if dark else LINE)
    value_h = value_size * 1.45 / 72
    text_box(slide, x + 0.18, y + 0.15, w - 0.36, value_h, [value], size=value_size,
             color=value_color, font=HEAD_FONT, autofit=False, line_spacing=1.0)
    text_box(slide, x + 0.18, y + 0.2 + value_h, w - 0.36, h - value_h - 0.3, [label], size=9.5,
             color=PAPER_DIM if dark else INK_SOFT, line_spacing=1.1)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def new_slide(prs, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY if dark else PAPER
    return s


# ---------------------------------------------------------------- slides
def slide_title(prs, n, total):
    s = new_slide(prs, dark=True)
    oval(s, 8.6, -3.2, 7.0, 7.0, NAVY_GLOW)                        # teal-ish glow, top right
    oval(s, -2.6, 5.2, 5.2, 5.2, RGBColor(0x2A, 0x22, 0x3A))       # coral-ish glow, bottom left
    # abstract "axon": thin steady line vs. thick swelling line
    hline(s, 0, 5.85, 13.333, LINE_DARK, 1.0)
    oval(s, 6.4, 5.55, 3.2, 1.0, RGBColor(0x1B, 0x2E, 0x4C))
    hline(s, 0, 6.05, 13.333, RGBColor(0x2B, 0x4A, 0x62), 6.0)
    chrome(s, n, total, "Rare neurogenetic disease", dark=True, right_label="OMIM #256850")

    kicker(s, 0.7, 1.55, 6, "Case study · gene to gene therapy")
    text_box(s, 0.7, 1.9, 7.6, 2.6, [
        {"text": "Giant Axonal", "size": 62},
        {"text": "__Neuropathy__", "size": 62},
    ], color=PAPER, font=HEAD_FONT, autofit=False, line_spacing=0.95, space_after=0)
    text_box(s, 0.7, 4.35, 6.4, 1.2, [
        "How the loss of one small protein, gigaxonin, causes nerve fibres to swell and die — "
        "and how one family's fight produced the world's first intrathecal gene therapy."
    ], size=15, color=PAPER_DIM, line_spacing=1.2)
    tags = [("GENE: GAN", TEAL_LIGHT, TEAL), ("CHR 16q23.2", PAPER_DIM, LINE_DARK),
            ("AUTOSOMAL RECESSIVE", PAPER_DIM, LINE_DARK), ("ONSET ≈ 3 YEARS", PAPER_DIM, LINE_DARK)]
    tx = 0.7
    for label, col, border in tags:
        w = 0.095 * len(label) + 0.45
        rect(s, tx, 6.35, w, 0.38, NAVY, line=border, radius=0.5)
        text_box(s, tx, 6.35, w, 0.38, [label], size=8.5, color=col, bold=True, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE, autofit=False)
        tx += w + 0.15

    agenda = ["The genetic root cause", "Gigaxonin and the broken pathway", "Phenotype: the giant axon",
              "Symptoms and clinical course", "Hannah Sames", "scAAV9/JeT-GAN gene therapy", "Results and legacy"]
    ay = 1.95
    for i, item in enumerate(agenda, 1):
        hline(s, 8.6, ay, 4.1, LINE_DARK, 0.75)
        text_box(s, 8.6, ay + 0.08, 0.6, 0.35, [f"{i:02d}"], size=13, color=TEAL, font=HEAD_FONT, autofit=False)
        text_box(s, 9.2, ay + 0.1, 3.5, 0.35, [item], size=12, color=PAPER, autofit=False)
        ay += 0.5
    notes(s, "Opening. Giant Axonal Neuropathy (GAN) is an ultra-rare, autosomal recessive neurodegenerative "
             "disease of childhood (OMIM #256850). The talk moves from the gene, to the cell, to the patient, "
             "to the first-in-human intrathecal gene therapy.")


def slide_genetics(prs, n, total):
    s = new_slide(prs)
    chrome(s, n, total, "01 · Genetic root cause")
    kicker(s, 0.7, 0.7, 6, "The root cause")
    heading(s, 0.7, 1.0, 6.4, "One gene, two broken copies")

    # spec card
    rect(s, 0.7, 1.95, 6.3, 1.85, TEAL_SOFT)
    rows = [("GENE", "**GAN** — chromosome 16, band 16q23.2"),
            ("PROTEIN", "**Gigaxonin** — 597 aa, BTB/Kelch family"),
            ("FUNCTION", "Substrate adaptor for a **Cul3 E3 ubiquitin ligase**"),
            ("MECHANISM", "Loss of function — **>100** variants (missense, nonsense, frameshift, deletions)")]
    ry = 2.08
    for label, val in rows:
        text_box(s, 0.9, ry, 1.2, 0.4, [label], size=8.5, color=TEAL, bold=True, autofit=False)
        text_box(s, 2.1, ry - 0.03, 4.8, 0.42, [val], size=12, color=INK, autofit=False)
        ry += 0.43

    bullets(s, 0.7, 4.0, 6.3, 1.7, [
        "**Autosomal recessive:** disease appears only when both inherited copies of __GAN__ are defective.",
        "**Carrier parents are healthy** — one working copy is enough. Each pregnancy carries a 25% risk of an affected child.",
        "Ultra-rare: prevalence below **1 in 1,000,000**; more frequent in consanguineous families.",
    ], size=12)

    _, (px, py, pw, ph) = picture_fit(s, ASSETS / "locus.jpg", 1.7, 5.65, 4.3, 1.1, align="left")
    caption(s, 0.7, 6.8, 6.3, "Mutation locus", "Long arm of chromosome 16, band q23.2 — home of the GAN gene")

    _, (px, py, pw, ph) = picture_fit(s, ASSETS / "inheritance.jpg", 7.5, 1.1, 5.2, 5.2)
    caption(s, px, py + ph + 0.1, pw, "Inheritance pattern", "Two carriers → 1 in 4 affected, 2 in 4 carriers, 1 in 4 unaffected")
    notes(s, "GAN is caused by biallelic loss-of-function variants in the GAN gene (16q23.2) encoding gigaxonin, "
             "a BTB-Kelch substrate adaptor for the Cul3 E3 ubiquitin ligase. Inheritance is autosomal recessive: "
             "carrier parents are unaffected; each pregnancy has a 25% risk. Prevalence is estimated below 1 per million.")


def slide_mechanism(prs, n, total):
    s = new_slide(prs)
    chrome(s, n, total, "02 · Molecular mechanism")
    kicker(s, 0.7, 0.7, 7, "What gigaxonin does — and what fails without it")
    heading(s, 0.7, 1.0, 12.0, "The cell's filament recycler goes missing")

    bullets(s, 0.7, 1.95, 5.9, 2.9, [
        "**Normal role:** gigaxonin tags **intermediate filament (IF)** proteins — neurofilaments, vimentin, peripherin, GFAP, keratins — for degradation by the proteasome.",
        "**Without gigaxonin:** IFs are never cleared. They pile up into dense, disorganised aggregates that choke axonal transport.",
        "**Second hit — Sonic hedgehog (Shh) signalling:** gigaxonin also degrades the receptor **Patched (Ptch)**. When Ptch accumulates, it keeps **Smoothened (Smo)** switched off, so the Shh pathway — key for neural development and repair — stays **OFF**.",
    ], size=12)

    _, (px, py, pw, ph) = picture_fit(s, ASSETS / "filament.jpg", 0.7, 4.75, 2.1, 2.1, align="left")
    caption(s, 2.95, 5.55, 3.6, "Intermediate filament aggregate", "Patient cell — the hallmark of gigaxonin loss (arrow), N = nucleus")

    _, (px, py, pw, ph) = picture_fit(s, ASSETS / "pathway.jpg", 7.0, 1.9, 5.7, 4.5)
    caption(s, px, py + ph + 0.1, pw, "Affected pathway", "With gigaxonin → Shh ON  ·  Without gigaxonin → Shh OFF")
    notes(s, "Gigaxonin controls turnover of all intermediate-filament classes. Loss leads to IF aggregation in neurons, "
             "glia, fibroblasts and hair. Gigaxonin also mediates degradation of Patched; without it Patched persists and "
             "Smoothened remains inhibited, silencing Sonic hedgehog signalling.")


def slide_phenotype(prs, n, total):
    s = new_slide(prs)
    chrome(s, n, total, "03 · Phenotype")
    kicker(s, 0.7, 0.7, 6, "The phenotype, cell by cell", CORAL)
    heading(s, 0.7, 1.0, 7.0, "Axons that swell into giants")

    bullets(s, 0.7, 1.95, 7.0, 2.9, [
        "**Focal axonal swellings** several times the normal diameter, packed with tangled neurofilaments — the \"giant axons\" that name the disease.",
        "**Myelin thins and degenerates** over each swelling; conduction slows, then fails.",
        "**Both nervous systems are hit:** peripheral sensory and motor nerves (PNS) and the cerebellum, brainstem and long white-matter tracts (CNS).",
        "**Not only neurons:** astrocytes fill with GFAP aggregates (Rosenthal fibres); disordered keratin filaments give the characteristic **tightly curled, \"kinky\" hair**.",
    ], size=12, dot=CORAL)

    cards = [("3–5×", "typical diameter of a swollen \"giant\" axon vs. a healthy one", CORAL),
             ("PNS + CNS", "a length-dependent axonopathy that later becomes a whole-brain disease", INK),
             ("MRI", "white-matter signal change in cerebellum and brainstem supports diagnosis", TEAL)]
    cx = 0.7
    for value, label, col in cards:
        stat_card(s, cx, 5.0, 2.2, 1.85, value, label, value_color=col, value_size=22)
        cx += 2.4

    _, (px, py, pw, ph) = picture_fit(s, ASSETS / "neuron.jpg", 8.1, 0.65, 4.6, 6.0)
    caption(s, px, py + ph + 0.08, pw, "Normal vs GAN", "Neurofilament accumulation → axonal swelling → myelin loss, in CNS and PNS")
    notes(s, "Histology: enlarged axons filled with densely packed neurofilaments, thin or absent myelin. "
             "Rosenthal fibres (GFAP aggregates in astrocytes) are common. Brain MRI shows cerebellar and "
             "brainstem white-matter signal change. Kinky hair reflects keratin IF disorganisation.")


def slide_symptoms(prs, n, total):
    s = new_slide(prs)
    chrome(s, n, total, "04 · Symptoms")
    kicker(s, 0.7, 0.9, 6, "Symptoms and clinical course", CORAL)
    heading(s, 0.7, 1.2, 11.5, "A slow, relentless loss — from the feet upward")
    text_box(s, 0.7, 2.0, 9.5, 0.8, [
        "GAN typically presents in a toddler who walks \"clumsily\", and progresses over two decades "
        "from peripheral weakness to whole-brain involvement."], size=14, color=INK_SOFT)

    stages = [
        ("ONSET · ≈ 3 YEARS", "The clumsy walker",
         "Unsteady gait, frequent falls, weakness in feet and hands, high-arched feet (pes cavus), diminished reflexes. "
         "Tightly curled hair unlike either parent is a visible clue.", TEAL),
        ("CHILDHOOD · 5–10 YEARS", "Peripheral neuropathy",
         "Loss of vibration and position sense, absent reflexes, muscle wasting below knees and elbows. "
         "Nerve biopsy shows giant axons; genetic testing confirms GAN variants.", GOLD_DARK),
        ("ADOLESCENCE", "Central nervous system",
         "Wheelchair dependence, cerebellar ataxia, nystagmus, slurred speech, scoliosis, seizures and cognitive "
         "decline as brain white matter degenerates.", CORAL),
        ("EARLY ADULTHOOD", "Late stage",
         "Swallowing and breathing muscles weaken; most patients die of respiratory failure in the 2nd–3rd decade. "
         "Until 2015 there was **no treatment beyond supportive care**.", INK_SOFT),
    ]
    hline(s, 0.7, 3.25, 11.9, LINE, 2.0)
    col_w, gap, x = 2.75, 0.3, 0.7
    for when, title, body, col in stages:
        oval(s, x - 0.02, 3.13, 0.24, 0.24, PAPER)
        ring = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.02), Inches(3.13), Inches(0.24), Inches(0.24))
        ring.fill.solid(); ring.fill.fore_color.rgb = PAPER; ring.line.color.rgb = col; ring.line.width = Pt(2.5); ring.shadow.inherit = False
        text_box(s, x, 3.55, col_w, 0.3, [when], size=8.5, color=col, bold=True, autofit=False)
        text_box(s, x, 3.85, col_w, 0.45, [title], size=15, color=INK, bold=True, autofit=False)
        text_box(s, x, 4.35, col_w, 2.0, [body], size=11.5, color=INK_SOFT, line_spacing=1.15)
        x += col_w + gap
    text_box(s, 0.7, 6.55, 11.9, 0.4, [
        "Diagnosis: clinical picture + nerve biopsy (giant axons) + brain MRI + GAN sequencing. "
        "Motor decline is tracked with the 32-item Motor Function Measure (MFM-32)."], size=9.5, color=INK_SOFT)
    notes(s, "Classic GAN begins before age 3 with gait disturbance and distal weakness, evolves into a severe "
             "sensorimotor neuropathy, then adds CNS signs (ataxia, nystagmus, dysarthria, seizures, cognitive decline). "
             "Most patients are wheelchair-bound in the second decade and die of respiratory failure by the third.")


def slide_hannah(prs, n, total):
    s = new_slide(prs, dark=True)
    oval(s, 9.8, 4.6, 6.0, 6.0, RGBColor(0x1E, 0x2A, 0x3E))
    chrome(s, n, total, "05 · One patient", dark=True)

    _, (px, py, pw, ph) = picture_fit(s, ASSETS / "hannah.jpg", 0.7, 1.6, 5.3, 4.4)
    badge_w = 4.1
    rect(s, px + 0.2, py + ph - 0.85, badge_w, 0.65, NAVY, radius=0.12)
    text_box(s, px + 0.35, py + ph - 0.8, badge_w - 0.3, 0.6, [
        {"text": "HANNAH SAMES", "size": 8.5, "color": TEAL_LIGHT, "bold": True, "space_after": 1},
        {"text": "Rexford, New York · diagnosed 24 March 2008, aged 4", "size": 9.5, "color": PAPER},
    ], autofit=False)

    kicker(s, 6.6, 1.05, 6, "The case that changed the field", GOLD)
    text_box(s, 6.6, 1.35, 6.2, 0.8, ["\"We decided to __fight__.\""], size=30, color=PAPER, font=HEAD_FONT, autofit=False)
    bullets(s, 6.6, 2.35, 6.1, 3.1, [
        "**2007:** at age 3 Hannah's walking and balance lag behind; doctors call it growing pains. Nearly a year of tests ends with a GAN diagnosis in **March 2008**.",
        "**The hardest genotype:** two deletion mutations — Hannah makes **no gigaxonin at all** (CRIM-negative), so her immune system had never \"seen\" the protein.",
        "**Hannah's Hope Fund** is founded within months by parents Lori and Matt Sames. Their first symposium (Boston, Aug 2008) gathers 20 experts, who conclude: __gene therapy is the only path fast enough__.",
        "The fund raises **>$8 million**, pays for all pre-clinical work and a natural-history study, and recruits a young UNC post-doc, **Steven Gray**, to build the vector.",
    ], size=11.5, color=PAPER_DIM)
    for r in s.shapes[-1].text_frame.paragraphs:
        for run in r.runs:
            if run.font.bold:
                run.font.color.rgb = PAPER
    rect(s, 6.6, 5.6, 0.04, 1.0, TEAL, rounded=False)
    text_box(s, 6.8, 5.55, 5.9, 1.2, [
        {"text": "Hannah's parents found themselves in the unimaginable situation of having funded a clinical trial that might have to exclude their own child.",
         "size": 13, "color": PAPER, "font": HEAD_FONT, "space_after": 4},
        {"text": "DNA SCIENCE, JULY 2016", "size": 8, "color": PAPER_DIM, "bold": True},
    ], line_spacing=1.15)
    notes(s, "Hannah Sames (born 2004, Rexford NY) was diagnosed on 24 March 2008. Her two deletion mutations mean she "
             "produces no gigaxonin (CRIM-negative). Lori and Matt Sames founded Hannah's Hope Fund, convened the first GAN "
             "symposium in August 2008, raised over $8 million and funded Steven Gray's vector development at UNC.")


def slide_therapy(prs, n, total):
    s = new_slide(prs)
    chrome(s, n, total, "06 · Treatment")
    kicker(s, 0.7, 0.62, 8, "The genetic treatment · scAAV9/JeT-GAN")
    heading(s, 0.7, 0.9, 11.5, "A working gene, delivered into the spinal fluid")
    text_box(s, 0.7, 1.65, 11.0, 0.7, [
        "The first-in-human trial of an AAV gene therapy given **intrathecally** — into the cerebrospinal fluid — "
        "opened at the NIH Clinical Center in 2015 (NCT02362438, led by Carsten Bönnemann, NINDS)."], size=13, color=INK_SOFT)

    steps = [
        ("1", "Build the vector", "**scAAV9**: a self-complementary adeno-associated virus 9 shell that can enter neurons and spread through spinal cord and brain. It carries no viral genes and does not integrate."),
        ("2", "Load the cargo", "A **codon-optimised human** __GAN__ transgene driven by **JeT**, a tiny synthetic promoter — small enough to fit the compact self-complementary genome."),
        ("3", "Single lumbar injection", "One dose — up to **3.5 × 10¹⁴** vector genomes — is injected into the CSF by lumbar puncture, bathing motor neurons and dorsal root ganglia directly and bypassing the blood–brain barrier."),
        ("4", "Restore gigaxonin", "Transduced neurons begin making gigaxonin within ~2 weeks, clearing filament aggregates and letting damaged sensory nerves **regenerate**."),
    ]
    cw, gap, x = 2.85, 0.17, 0.7
    for num, title, body in steps:
        rect(s, x, 2.5, cw, 2.55, WHITE, line=LINE)
        text_box(s, x + 0.2, 2.6, 0.6, 0.5, [num], size=22, color=TEAL, font=HEAD_FONT, autofit=False)
        text_box(s, x + 0.2, 3.05, cw - 0.4, 0.35, [title], size=13, color=INK, bold=True, autofit=False)
        text_box(s, x + 0.2, 3.4, cw - 0.4, 1.6, [body], size=10.5, color=INK_SOFT, line_spacing=1.12)
        if num != "4":
            text_box(s, x + cw - 0.02, 3.55, gap + 0.04, 0.3, ["→"], size=12, color=TEAL, bold=True,
                     align=PP_ALIGN.CENTER, autofit=False)
        x += cw + gap

    stats = [("May 2015", "first human ever to receive a therapeutic gene via the spinal fluid — a GAN patient", TEAL),
             ("21 Jul 2016", "Hannah, aged 12, treated as the 5th child (≈120 trillion vector genomes)", GOLD_DARK),
             ("CRIM-neg", "as the first patient making no gigaxonin, Hannah needed a dedicated immunosuppression protocol", CORAL),
             ("UNC → UTSW", "vector designed by Steven Gray in Jude Samulski's UNC Gene Therapy Center, funded by Hannah's Hope Fund", INK)]
    x = 0.7
    for value, label, col in stats:
        stat_card(s, x, 5.25, cw, 1.6, value, label, value_color=col, value_size=20)
        x += cw + gap
    notes(s, "scAAV9/JeT-GAN: self-complementary AAV9 capsid, codon-optimised human GAN cDNA, minimal synthetic JeT promoter. "
             "Developed by Steven Gray (UNC, now UT Southwestern) with funding from Hannah's Hope Fund. Delivered once by "
             "lumbar intrathecal injection at the NIH Clinical Center (NCT02362438, PI Carsten Bönnemann). First patient dosed "
             "May 2015; Hannah treated 21 July 2016 at 1.2 × 10^14 vg with an added immunosuppression protocol because she is CRIM-negative.")


def slide_results(prs, n, total):
    s = new_slide(prs, dark=True)
    oval(s, -2.8, -3.0, 6.5, 6.5, NAVY_GLOW)
    chrome(s, n, total, "07 · Results & legacy", dark=True)
    kicker(s, 0.7, 0.7, 8, "What the trial showed — NEJM, March 2024")
    heading(s, 0.7, 1.0, 11.5, "Bending the curve of an \"untreatable\" disease", dark=True)

    stats = [("14", "children aged 6–14 treated at four dose levels, followed for a median of 68.7 months"),
             ("1", "serious adverse event (fever) possibly related to treatment; CSF inflammation was managed with steroids"),
             ("−7.2 → −1.9", "MFM-32 points lost per year before vs. after the 1.8 × 10¹⁴ vg dose (slope slowed by 5.3 points)"),
             ("99%", "posterior probability that motor decline slowed at that dose — above the prespecified 95% bar"),
             ("6 / 14", "regained, stabilised or newly recordable sensory nerve responses 6–24 months after dosing"),
             ("3–48×", "more regenerating nerve-fibre clusters in sensory-nerve biopsies one year after treatment")]
    cw, ch, gap = 2.0, 1.85, 0.15
    for i, (value, label) in enumerate(stats):
        col, row = i % 3, i // 3
        x, y = 0.7 + col * (cw + gap), 2.0 + row * (ch + gap)
        vs = 18 if len(value) > 6 else 26
        stat_card(s, x, y, cw, ch, value, label, value_color=PAPER, dark=True, value_size=vs)

    bullets(s, 7.4, 1.95, 5.3, 2.6, [
        "**Not a cure — yet.** Damage already done is not reversed, and motor nerves kept declining; benefit is largest when children are treated **young and early**.",
        "**A template for the field:** intrathecal AAV9 delivery pioneered here is now used in trials for many other childhood neurological diseases.",
        "**Next steps:** the programme was licensed to industry (Taysha, TSHA-120); Hannah's Hope Fund funds a second injection site and research into treating the autonomic nervous system.",
    ], size=11.5, color=PAPER_DIM)
    for p in s.shapes[-1].text_frame.paragraphs:
        for run in p.runs:
            if run.font.bold:
                run.font.color.rgb = PAPER
    rect(s, 7.4, 4.75, 0.04, 1.15, TEAL, rounded=False)
    text_box(s, 7.6, 4.7, 5.1, 1.3, [
        {"text": "One family with no medical background turned a diagnosis into a first-in-human therapy — proof of what patient-driven science can do.",
         "size": 13, "color": PAPER, "font": HEAD_FONT, "space_after": 4},
        {"text": "TAKE-HOME MESSAGE", "size": 8, "color": PAPER_DIM, "bold": True},
    ], line_spacing=1.15)
    text_box(s, 7.4, 6.2, 5.3, 0.6, [
        "Bharucha-Goebel DX et al. __Intrathecal Gene Therapy for Giant Axonal Neuropathy.__ N Engl J Med 2024;390:1092-1104."],
        size=9, color=PAPER_DIM)
    notes(s, "Bharucha-Goebel et al., NEJM 2024: 14 children, doses 3.5×10^13 to 3.5×10^14 vg, median follow-up 68.7 months. "
             "One serious adverse event possibly related (fever). Pre-treatment MFM-32 slope −7.17 points/year; at 1.8×10^14 vg "
             "the slope improved by 5.32 points (99% posterior probability of slowing, above the 95% threshold). Sensory nerve "
             "action potentials recovered or stabilised in 6 of 14; nerve biopsies showed 3–48× more regenerating clusters. "
             "Two older low-dose participants died of disease-related events. The programme was licensed to Taysha (TSHA-120).")


# ---------------------------------------------------------------- build
def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    builders = [slide_title, slide_genetics, slide_mechanism, slide_phenotype,
                slide_symptoms, slide_hannah, slide_therapy, slide_results]
    total = len(builders)
    for i, fn in enumerate(builders, 1):
        fn(prs, i, total)
    prs.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size // 1024} KB, {total} slides)")


if __name__ == "__main__":
    build()
