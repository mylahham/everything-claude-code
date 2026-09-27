#!/usr/bin/env python3
"""Build an editable, Canva-importable PPTX of the visual-first roll-up banner.

Mirrors giant-axonal-neuropathy-banner-visual.html: one slide in the banner's
850 x 2000 mm proportions, every element a separate text box / shape / picture,
fonts from Canva's library (Fraunces + Manrope), 2x upscaled figures from
../assets/hires and the vector icons exported by export-icons.js (canva-assets/).

The PPTX format caps a slide at 56 in (1422 mm) per side, so the slide is emitted
at 70 % (595 x 1400 mm) - the largest in-spec size. Everything scales losslessly.

Usage:   python3 build_canva_pptx.py          # 595 x 1400 mm  (recommended for Canva)
         python3 build_canva_pptx.py --half   # 425 x 1000 mm
Requires: pip install python-pptx pillow
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

HERE = Path(__file__).parent
HIRES = HERE.parent / "assets" / "hires"
ICONS = HERE / "canva-assets"

S = 0.5 if "--half" in sys.argv else 0.7
W_MM, H_MM = 850, 2000
OUT = HERE / f"giant-axonal-neuropathy-banner-visual-canva-{W_MM * S:.0f}x{H_MM * S:.0f}mm.pptx"

# ------------------------------------------------------------------ theme
PAPER, PAPER_DEEP = RGBColor(0xF6, 0xF3, 0xEC), RGBColor(0xEC, 0xE7, 0xDB)
INK, INK_SOFT = RGBColor(0x14, 0x21, 0x3D), RGBColor(0x4A, 0x55, 0x70)
NAVY, NAVY_2, NAVY_GLOW = RGBColor(0x0F, 0x1B, 0x33), RGBColor(0x18, 0x27, 0x47), RGBColor(0x12, 0x30, 0x45)
TEAL, TEAL_SOFT, TEAL_LIGHT = RGBColor(0x1A, 0x9C, 0x8E), RGBColor(0xDD, 0xEF, 0xEC), RGBColor(0x7F, 0xE0, 0xD4)
CORAL, GOLD, GOLD_DARK = RGBColor(0xE4, 0x57, 0x2E), RGBColor(0xE9, 0xB4, 0x4C), RGBColor(0xB7, 0x80, 0x1E)
WHITE, LINE, LINE_DARK = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xDC, 0xD8, 0xCE), RGBColor(0x2A, 0x3A, 0x5C)
PAPER_DIM, GHOST = RGBColor(0xB8, 0xBC, 0xC8), RGBColor(0xE2, 0xDF, 0xD9)

HEAD, BODY = "Fraunces", "Manrope"

T_TITLE, T_H2, T_H3, T_BODY, T_SMALL, T_KICK, T_STAT = 132, 62, 30, 26, 17, 18, 74

TOKEN = re.compile(r"(\*\*.+?\*\*|__.+?__)")


# ------------------------------------------------------------------ helpers
def mm(v):
    return Emu(int(round(v * S * 36000)))


def pt(v):
    return Pt(v * S)


def add_runs(p, text, size, color, font=BODY, bold=False, italic=False):
    for part in TOKEN.split(text):
        if not part:
            continue
        b, i, f = bold, italic, font
        if part.startswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("__"):
            part, i, f = part[2:-2], True, HEAD
        r = p.add_run()
        r.text = part
        r.font.name, r.font.size, r.font.bold, r.font.italic = f, pt(size), b, i
        r.font.color.rgb = color


def text(slide, x, y, w, h, paras, *, size=T_BODY, color=INK, font=BODY, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ls=1.15, after=3):
    box = slide.shapes.add_textbox(mm(x), mm(y), mm(w), mm(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    for i, item in enumerate(paras):
        spec = item if isinstance(item, dict) else {"text": item}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = spec.get("align", align)
        p.line_spacing = spec.get("ls", ls)
        p.space_after = pt(spec.get("after", after))
        add_runs(p, spec["text"], spec.get("size", size), spec.get("color", color),
                 spec.get("font", font), spec.get("bold", bold), spec.get("italic", italic))
    return box


def bullets(slide, x, y, w, h, items, *, size=T_BODY, color=INK_SOFT, dot=TEAL, gap=5, ls=1.2, strong=INK):
    box = slide.shapes.add_textbox(mm(x), mm(y), mm(w), mm(h))
    tf = box.text_frame
    tf.word_wrap = True
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = ls
        p.space_after = pt(gap)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(mm(10))))
        pPr.set("indent", str(-int(mm(10))))
        r = p.add_run()
        r.text = "\u25cf\u2002"
        r.font.name, r.font.size = BODY, pt(size * 0.6)
        r.font.color.rgb = dot
        add_runs(p, item, size, color)
    for p in tf.paragraphs:
        for r in p.runs:
            if r.font.bold and r.font.color.rgb == color:
                r.font.color.rgb = strong
    return box


def rect(slide, x, y, w, h, fill, *, line=None, rounded=True, radius=0.05, lw=1.2):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                 mm(x), mm(y), mm(w), mm(h))
    if rounded:
        shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = pt(lw)
    shp.shadow.inherit = False
    return shp


def oval(slide, x, y, w, h, fill, *, line=None, lw=2.0):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, mm(x), mm(y), mm(w), mm(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = pt(lw)
    shp.shadow.inherit = False
    return shp


def picture(slide, path, x, y, w=None, h=None):
    """Place an image by width OR height (mm), keeping aspect ratio. Returns (x, y, w, h) in mm."""
    with Image.open(path) as im:
        iw, ih = im.size
    if w is None:
        w = h * iw / ih
    elif h is None:
        h = w * ih / iw
    slide.shapes.add_picture(str(path), mm(x), mm(y), mm(w), mm(h))
    return x, y, w, h


def caption(slide, x, y, w, label, body, *, dark=False):
    box = text(slide, x, y, w, 16, [""], size=T_SMALL, ls=1.25)
    p = box.text_frame.paragraphs[0]
    add_runs(p, label.upper() + "   ", T_SMALL * 0.82, PAPER if dark else INK, bold=True)
    add_runs(p, body, T_SMALL, PAPER_DIM if dark else INK_SOFT)
    return box


def kicker(slide, x, y, w, label, color=TEAL):
    rect(slide, x, y + 3.4, 18, 1.4, color, rounded=False)
    text(slide, x + 24, y, w, 10, [label.upper()], size=T_KICK, color=color, bold=True)


def head(slide, y, kick, title, *, kcolor=TEAL, tcolor=INK, tsize=T_H2):
    kicker(slide, 38, y + 16, 720, kick, kcolor)
    text(slide, 38, y + 27, 774, 26, [title], size=tsize, color=tcolor, font=HEAD, ls=1.0)


def section_chrome(slide, y, number, *, dark=False, color=TEAL):
    oval(slide, 6, y + 22, 26, 26, NAVY if dark else PAPER, line=color, lw=9)
    text(slide, 6, y + 22, 26, 26, [str(number)], size=22, color=PAPER if dark else INK, font=HEAD,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def stat(slide, x, y, w, h, value, label, *, vcolor=INK, vsize=T_STAT, lsize=T_SMALL, dark=False, pad=10):
    rect(slide, x, y, w, h, NAVY_2 if dark else WHITE, line=None if dark else LINE)
    vh = vsize * 0.353 * 1.15
    text(slide, x + pad, y + pad, w - 2 * pad, vh, [value], size=vsize, color=vcolor, font=HEAD, ls=1.0)
    text(slide, x + pad, y + pad + vh + 3, w - 2 * pad, h - vh - pad - 6, [label], size=lsize,
         color=PAPER_DIM if dark else INK_SOFT, ls=1.25)


# ------------------------------------------------------------------ build
def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = mm(W_MM), mm(H_MM)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = PAPER
    X, CW = 38, W_MM - 2 * 38

    # ---------------- HEADER 0-170
    rect(s, 0, 0, W_MM, 170, NAVY, rounded=False)
    oval(s, 640, -230, 380, 380, NAVY_GLOW)
    oval(s, -150, -40, 200, 200, RGBColor(0x2A, 0x22, 0x3A))
    rect(s, 0, 152, W_MM, 1.2, LINE_DARK, rounded=False)
    oval(s, 380, 142, 220, 30, RGBColor(0x1B, 0x2E, 0x4C))
    rect(s, 0, 158, W_MM, 4, RGBColor(0x2B, 0x4A, 0x62), rounded=False)
    kicker(s, X, 22, 600, "Rare neurogenetic disease · OMIM #256850")
    text(s, X, 34, 780, 52, ["Giant Axonal __Neuropathy__"], size=T_TITLE, color=PAPER, font=HEAD, ls=1.0)
    text(s, X, 92, 700, 16, ["From a missing protein to the first gene therapy delivered into the spinal fluid."],
         size=28, color=PAPER_DIM, ls=1.15)
    tx = X
    for label, col, border in [("Gene: GAN", TEAL_LIGHT, TEAL), ("Chromosome 16q23.2", PAPER_DIM, LINE_DARK),
                               ("Autosomal recessive", PAPER_DIM, LINE_DARK), ("Onset ≈ 3 years", PAPER_DIM, LINE_DARK),
                               ("< 1 in 1,000,000", PAPER_DIM, LINE_DARK)]:
        w = 4.6 * len(label) + 18
        rect(s, tx, 118, w, 14, NAVY, line=border, radius=0.5)
        text(s, tx, 118, w, 14, [label.upper()], size=T_SMALL, color=col, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tx += w + 6

    # ---------------- 1 · ONE GENE, TWO BROKEN COPIES 170-500
    y = 170
    rect(s, 17.5, y, 3, 330, TEAL, rounded=False)
    section_chrome(s, y, 1)
    head(s, y, "The genetic root cause", "One gene, two broken copies")
    lw_col = CW - 16 - 300
    fw = (lw_col - 16) / 3
    for i, (k, v, sub) in enumerate([("Gene", "__GAN__", "chromosome 16 · band q23.2"),
                                     ("Protein", "Gigaxonin", "597 aa · BTB / Kelch family"),
                                     ("Variants", "> 100", "all loss-of-function")]):
        fx = X + i * (fw + 8)
        rect(s, fx, y + 62, fw, 46, TEAL_SOFT)
        text(s, fx + 10, y + 68, fw - 20, 8, [k.upper()], size=14, color=TEAL, bold=True)
        text(s, fx + 10, y + 76, fw - 20, 16, [v], size=34, color=INK, font=HEAD, ls=1.0)
        text(s, fx + 10, y + 93, fw - 20, 10, [sub], size=16, color=INK_SOFT)
    _, _, lw_, lh_ = picture(s, HIRES / "locus.png", X, y + 122, h=150)
    caption(s, X, y + 278, lw_col, "Mutation locus", "Long arm of chromosome 16, band q23.2")
    ix, iy, iw, ih = picture(s, HIRES / "inheritance.png", X + lw_col + 16, y + 62, h=238)
    caption(s, ix, iy + ih + 6, 300, "Two healthy carriers", "each child: 1 in 4 affected · 2 in 4 carriers · 1 in 4 unaffected")

    # ---------------- 2 · THE GIGAXONIN PATHWAY 500-1010
    y = 500
    rect(s, 0, y, W_MM, 510, PAPER_DEEP, rounded=False)
    rect(s, 17.5, y, 3, 510, TEAL, rounded=False)
    section_chrome(s, y, 2)
    head(s, y, "The gigaxonin pathway — how one missing protein damages the nerve", "The cell's filament recycler goes missing")
    steps = [("gene", TEAL, "__GAN__ mutated", "Both copies carry a loss-of-function variant"),
             ("missing", TEAL, "No gigaxonin", "The adaptor of the Cul3 E3 ubiquitin ligase is absent"),
             ("tag", GOLD_DARK, "Filaments not tagged", "Neurofilaments, vimentin, GFAP and keratin are never ubiquitinated"),
             ("pile", GOLD_DARK, "Aggregates pile up", "The proteasome cannot clear them; dense tangles form"),
             ("axon", CORAL, "Axon swells", "Focal \"giant\" swellings; myelin thins over them"),
             ("signal", CORAL, "Signal fails", "Transport and conduction break down in nerves and brain")]
    colw = (CW - 5 * 8) / 6
    cy = y + 62
    centers = [X + i * (colw + 8) + colw / 2 for i in range(6)]
    for (x0, x1, col) in [(centers[0], centers[2], TEAL), (centers[2], centers[4], GOLD_DARK), (centers[4], centers[5], CORAL)]:
        rect(s, x0, cy + 32, x1 - x0, 2, col, rounded=False)
    for i, (icon, col, title, body) in enumerate(steps):
        cx = centers[i]
        picture(s, ICONS / f"cascade-{icon}.png", cx - 33, cy, w=66)
        oval(s, cx + 18, cy - 4, 16, 16, col)
        text(s, cx + 18, cy - 4, 16, 16, [str(i + 1)], size=14, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, cx - colw / 2, cy + 72, colw, 14, [title], size=25, color=INK, bold=True, align=PP_ALIGN.CENTER, ls=1.05)
        text(s, cx - colw / 2, cy + 87, colw, 30, [body], size=18, color=INK_SOFT, align=PP_ALIGN.CENTER, ls=1.22)
        if i < 5:
            text(s, cx + colw / 2 - 4, cy + 14, 16, 22, ["›"], size=44, color=INK_SOFT, bold=True, align=PP_ALIGN.CENTER)
    iy = cy + 127
    fx, fy, fw_, fh = picture(s, HIRES / "filament.png", X, iy, h=276)
    caption(s, fx, fy + fh + 5, fw_, "Step 4", "Intermediate-filament aggregate in a patient cell — the hallmark of gigaxonin loss")
    nx, ny, nw, nh = picture(s, HIRES / "neuron.png", fx + fw_ + 16, iy, h=276)
    caption(s, nx, ny + nh + 5, nw + 10, "Step 5", "Normal vs GAN neuron: accumulation → swelling → myelin loss")
    sx = nx + nw + 16
    sw = W_MM - 38 - sx
    ch = (276 - 2 * 10) / 3
    for i, (v, l, c, vs) in enumerate([("3–5×", "diameter of a swollen \"giant\" axon compared with a healthy one", CORAL, T_STAT),
                                        ("PNS + CNS", "peripheral nerves first, then cerebellum, brainstem and white matter", INK, 52),
                                        ("Curly hair", "disordered keratin filaments — a visible clue at diagnosis", GOLD_DARK, 52)]):
        stat(s, sx, iy + i * (ch + 10), sw, ch, v, l, vcolor=c, vsize=vs)

    # ---------------- 3 · SYMPTOMS 1010-1280
    y = 1010
    rect(s, 17.5, y, 3, 270, GOLD_DARK, rounded=False)
    section_chrome(s, y, 3, color=GOLD_DARK)
    head(s, y, "Symptoms and clinical course", "A slow loss — from the feet upward", kcolor=CORAL)
    ty = y + 76
    colw = (CW - 3 * 12) / 4
    for (x0, x1, col) in [(X, X + colw + 12, TEAL), (X + colw + 12, X + 2 * (colw + 12), GOLD_DARK),
                          (X + 2 * (colw + 12), X + 3 * (colw + 12), CORAL), (X + 3 * (colw + 12), X + CW, INK_SOFT)]:
        rect(s, x0, ty + 6, x1 - x0, 1.8, col, rounded=False)
    stages = [("teal", TEAL, "≈ 3 years", "Clumsy walker", ["Unsteady gait, falls", "Weak feet, high arches", "Tightly curled hair"]),
              ("gold", GOLD_DARK, "5–10 years", "Peripheral nerves", ["Lost reflexes and sensation", "Muscle wasting", "Biopsy: giant axons"]),
              ("coral", CORAL, "Adolescence", "Brain involved", ["Wheelchair, ataxia", "Slurred speech, nystagmus", "Scoliosis, seizures"]),
              ("ink", INK_SOFT, "20s", "Late stage", ["Swallowing, breathing fail", "Respiratory failure", "**No treatment until 2015**"])]
    for i, (icon, col, when, title, items) in enumerate(stages):
        cx = X + i * (colw + 12)
        oval(s, cx, ty, 14, 14, PAPER, line=col, lw=7)
        picture(s, ICONS / f"body-{icon}.png", cx, ty + 20, w=66)
        tx = cx + 74
        tw = colw - 74
        text(s, tx, ty + 22, tw, 9, [when.upper()], size=T_SMALL, color=col, bold=True)
        text(s, tx, ty + 33, tw, 16, [title], size=31, color=INK, bold=True, ls=1.05)
        bullets(s, tx, ty + 52, tw, 70, items, size=23, color=INK_SOFT, dot=col, gap=5, ls=1.2)

    # ---------------- 4 · HANNAH 1280-1500 (dark)
    y = 1280
    rect(s, 0, y, W_MM, 220, NAVY, rounded=False)
    oval(s, 730, y + 60, 150, 150, RGBColor(0x1E, 0x2A, 0x3E))
    rect(s, 17.5, y, 3, 220, GOLD, rounded=False)
    section_chrome(s, y, 4, dark=True, color=GOLD)
    hx, hy, hw, hh = picture(s, HIRES / "hannah.png", X, y + 16, w=296)
    rect(s, hx + 8, hy + hh - 30, 150, 22, NAVY, radius=0.15)
    text(s, hx + 14, hy + hh - 27, 140, 18, [
        {"text": "HANNAH SAMES", "size": T_SMALL * 0.8, "color": TEAL_LIGHT, "bold": True, "after": 1},
        {"text": "Rexford, New York", "size": T_SMALL * 0.9, "color": PAPER},
    ])
    cx = X + 296 + 22
    cw = W_MM - 38 - cx
    kicker(s, cx, y + 36, 500, "The case that changed the field", GOLD)
    text(s, cx, y + 48, cw, 26, ["\"We decided to __fight__.\""], size=T_H2, color=PAPER, font=HEAD, ls=1.0)
    mw = (cw - 20) / 3
    for i, (yr, t, d, col) in enumerate([("2008", "Diagnosed at 4", "Two deletions — Hannah makes no gigaxonin at all", GOLD),
                                         ("2008", "Hannah's Hope Fund", "Founded by her parents; first GAN symposium the same year", TEAL_LIGHT),
                                         ("$8 M+", "Raised for a cure", "Funds Steven Gray (UNC) to build the gene-therapy vector", CORAL)]):
        mx = cx + i * (mw + 10)
        rect(s, mx, y + 92, mw, 2.2, col, rounded=False)
        text(s, mx, y + 100, mw, 20, [yr], size=48, color=col, font=HEAD, ls=1.0)
        text(s, mx, y + 122, mw, 12, [t], size=24, color=PAPER, bold=True)
        text(s, mx, y + 136, mw, 40, [d], size=18, color=PAPER_DIM, ls=1.25)

    # ---------------- 5 · GENE THERAPY 1500-1770
    y = 1500
    rect(s, 17.5, y, 3, 270, TEAL, rounded=False)
    section_chrome(s, y, 5)
    head(s, y, "The genetic treatment · scAAV9/JeT-GAN · NIH Clinical Center · NCT02362438", "A working gene, delivered into the spinal fluid")
    tsteps = [("capsid", "Build the vector", "scAAV9 shell: enters neurons, spreads through cord and brain"),
              ("cargo", "Load the gene", "Healthy human __GAN__ gene with a tiny JeT promoter"),
              ("inject", "One lumbar injection", "Into the CSF — bypassing the blood–brain barrier"),
              ("neuron", "Gigaxonin restored", "Neurons clear aggregates; sensory nerves regenerate")]
    colw = (CW - 3 * 12) / 4
    sy = y + 68
    for i, (icon, title, body) in enumerate(tsteps):
        cx = X + i * (colw + 12)
        rect(s, cx, sy, colw, 92, WHITE, line=LINE)
        picture(s, ICONS / f"therapy-{icon}.png", cx + 12, sy + 15, w=62)
        oval(s, cx + 10, sy - 7, 18, 18, TEAL)
        text(s, cx + 10, sy - 7, 18, 18, [str(i + 1)], size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, cx + 82, sy + 16, colw - 94, 26, [title], size=26, color=INK, bold=True, ls=1.05)
        text(s, cx + 82, sy + 44, colw - 94, 40, [body], size=19, color=INK_SOFT, ls=1.25)
        if i < 3:
            text(s, cx + colw - 3, sy + 38, 18, 16, ["→"], size=28, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    sw3 = (CW - 2 * 12) / 3
    for i, (v, l, c) in enumerate([("May 2015", "first person ever to receive a therapeutic gene through the spinal fluid", TEAL),
                                    ("21 July 2016", "Hannah, aged 12, becomes the 5th child treated", GOLD_DARK),
                                    ("1 dose", "up to 3.5 × 10¹⁴ vector genomes, given once", CORAL)]):
        stat(s, X + i * (sw3 + 12), sy + 92 + 16, sw3, 72, v, l, vcolor=c, vsize=56, lsize=19, pad=12)

    # ---------------- 6 · RESULTS 1770-1940 (dark)
    y = 1770
    rect(s, 0, y, W_MM, 230, NAVY, rounded=False)
    oval(s, -140, y + 10, 220, 220, NAVY_GLOW)
    rect(s, 17.5, y, 3, 170, CORAL, rounded=False)
    section_chrome(s, y, 6, dark=True, color=CORAL)
    lw_col = CW - 20 - 250
    kicker(s, X, y + 20, 600, "What the trial showed — NEJM, March 2024")
    text(s, X, y + 31, lw_col, 22, ["Bending the curve of an \"untreatable\" disease"], size=50, color=PAPER, font=HEAD, ls=1.0)
    rw = (lw_col - 3 * 12) / 4
    for i, (v, l) in enumerate([("14", "children treated, followed ≈ 6 years"), ("99%", "probability motor decline slowed at the best dose"),
                                ("6 / 14", "regained or stabilised sensory nerve responses"), ("1", "serious adverse event, managed with steroids")]):
        stat(s, X + i * (rw + 12), y + 64, rw, 76, v, l, vcolor=PAPER, vsize=T_STAT, dark=True)
    qx = X + lw_col + 20
    rect(s, qx, y + 60, 2.5, 60, TEAL, rounded=False)
    text(s, qx + 10, y + 58, 240, 70, ["Not a cure yet — but the first proof that an \"untreatable\" childhood neuropathy can be slowed."],
         size=30, color=PAPER, font=HEAD, ls=1.2)

    # ---------------- FOOTER 1940-2000
    y = 1940
    rect(s, 0, y, W_MM, 0.8, LINE_DARK, rounded=False)
    text(s, X, y + 12, 540, 40, [
        "Sources: Bharucha-Goebel DX et al. __Intrathecal Gene Therapy for Giant Axonal Neuropathy.__ N Engl J Med 2024;390:1092-1104 · "
        "ClinicalTrials.gov NCT02362438 · NIH News, March 2024 · Hannah's Hope Fund · The Children's Inn at NIH, \"Hannah's Story\" · "
        "DNA Science (PLOS), July 2016. Figures adapted from the source poster."], size=12, color=PAPER_DIM, ls=1.3)
    text(s, 600, y + 14, 212, 36, [
        {"text": "Giant Axonal Neuropathy", "size": 28, "color": PAPER, "font": HEAD, "align": PP_ALIGN.RIGHT, "after": 2},
        {"text": "FROM GENE TO GENE THERAPY", "size": 12, "color": TEAL_LIGHT, "bold": True, "align": PP_ALIGN.RIGHT},
    ])

    prs.save(OUT)
    print(f"wrote {OUT.name}  ({OUT.stat().st_size // 1024} KB)  slide {W_MM * S:.0f} x {H_MM * S:.0f} mm")


if __name__ == "__main__":
    build()
